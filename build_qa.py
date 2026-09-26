"""Update the v3 knowledge database from a reviewed Q&A DOCX.

Usage: python build_qa.py path/to/Uhamiaji_Tanzania_QA.docx
This deliberately does not edit embed.js or any widget asset.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
SOURCE = Path(sys.argv[1]).resolve()
DATA = ROOT / "data" / "knowledge.json"
OUTPUT = ROOT / "Mr_HamaHama_QA_Zote_Kwa_Kategoria.docx"

# Questions on the same issue replace older answers; unrelated old FAQs stay.
MATCH = {
    1:"visa.apply", 2:"visa.what", 3:"visa.on_arrival", 7:"visa.passport_req",
    8:"visa.entry", 9:"visa.zanzibar", 10:"visa.ordinary", 11:"visa.usa",
    12:"visa.multiple", 14:"visa.transit", 17:"visa.children", 27:"visa.time",
    28:"visa.track", 31:"visa.refund", 32:"visa.referral", 45:"visa.business",
    49:"mjue.visitor_work", 51:"residence.class_a", 52:"residence.class_b",
    53:"residence.class_c", 58:"pass.dependant", 59:"mjue.marriage",
    61:"residence.real_estate", 62:"residence.retiree", 67:"passport.apply",
    68:"passport.requirements", 70:"passport.minor", 71:"passport.abroad",
    72:"passport.fee", 73:"passport.payment", 74:"passport.validity",
    75:"passport.lost", 77:"passport.damaged", 81:"passport.etd",
    83:"passport.what", 84:"passport.service_req", 85:"citizenship.married",
    86:"citizenship.naturalisation", 89:"citizenship.dual",
    91:"citizenship.descent", 92:"citizenship.birth", 94:"visa.fee",
    96:"visa.gratis", 99:"passport.lost_fee", 117:"general.illegal",
    118:"mjue.refugees", 120:"passport.ci_ctd",
}
EXTRA_ALIASES = {67: ["ninaombaje pasipoti"]}


def classification(number):
    if number <= 44 or 93 <= number <= 97:
        return "visa", "VISA", "Visa"
    if 67 <= number <= 84 or 99 <= number <= 100 or number == 120:
        return "passport", "PASIPOTI NA HATI ZA SAFARI", "Passports and travel documents"
    if 85 <= number <= 92 or 109 <= number <= 112:
        return "citizenship", "URAIA", "Citizenship"
    if 113 <= number <= 116:
        return "general", "HUDUMA KWA WATEJA", "Customer services"
    if 117 <= number <= 119:
        return "mjue", "HIFADHI NA UHAMIAJI HARAMU", "Refugees and irregular migration"
    if number in (46, 48, 58, 101, 102, 105, 106):
        return "pass", "PASI NA FOMU", "Passes and forms"
    return "residence", "VIBALI VYA UKAAZI NA AJIRA", "Residence and employment permits"


def language(number):
    return "sw" if 67 <= number <= 92 or 99 <= number <= 100 or 109 <= number <= 120 else "en"


def answer_text(raw):
    return re.sub(r"\s*\[[^\[\]]+\]\s*$", "", raw).strip()


doc = Document(SOURCE)
parsed = []
for index, para in enumerate(doc.paragraphs):
    if para.style.name != "Heading 2":
        continue
    match = re.match(r"^(\d+)\.\s*(.+)$", para.text.strip())
    if not match:
        raise ValueError(f"Unexpected heading: {para.text!r}")
    number, question = int(match.group(1)), match.group(2).strip()
    raw_answer = doc.paragraphs[index + 1].text.strip()
    if not raw_answer:
        raise ValueError(f"Answer missing at {number}")
    parsed.append((number, question, raw_answer))
if [number for number, _, _ in parsed] != list(range(1, 121)):
    raise ValueError("Expected 120 sequential Q&A pairs")

kb = json.loads(DATA.read_text(encoding="utf-8"))
# Rebuilds remain idempotent: remove previous imports before matching.
kb["faqs"] = [f for f in kb["faqs"] if not f["id"].startswith("docx2026.")]
by_id = {f["id"]: f for f in kb["faqs"]}
if len(set(MATCH.values())) != len(MATCH):
    raise ValueError("One existing FAQ was mapped to multiple document questions")
if not set(MATCH.values()) <= set(by_id):
    raise ValueError("Mapped FAQ id absent from v3 database")

for number, question, raw_answer in parsed:
    topic, group_sw, group_en = classification(number)
    lang = language(number)
    other = "en" if lang == "sw" else "sw"
    text = answer_text(raw_answer)
    if number in MATCH:
        faq = by_id[MATCH[number]]
        faq["aliases"] = list(dict.fromkeys((faq.get("aliases") or []) + [question] + EXTRA_ALIASES.get(number, [])))
        faq["q"][lang] = question
    else:
        faq = {
            "id": f"docx2026.{number:03d}",
            "topic": topic,
            "group": {"sw": group_sw, "en": group_en},
            "q": {"sw": question, "en": question},
            "aliases": [question],
            "answer": {},
            "related": [],
        }
        kb["faqs"].append(faq)
    # Both keys are required by the current widget. Never silently reuse an
    # older potentially contradictory translation as the answer of record.
    faq["answer"] = {
        lang: {"text": text},
        other: {"text": ("English translation pending review. Original answer: " if lang == "sw" else "Tafsiri ya Kiswahili inahitaji uhakiki. Jibu la asili: ") + text},
    }
    faq["source_update"] = {"document": SOURCE.name, "question_number": number, "original_language": lang, "official_approval": False}

kb["version"] = "2026-09-26-qa"
DATA.write_text(json.dumps(kb, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

review = Document()
normal = review.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10)
review.add_heading("Mr. HamaHama - Maswali na Majibu Yote", 0)
review.add_paragraph(
    f"Jumla: {len(kb['faqs'])} Q&A. Nyaraka ya leo imeleta maswali 120; "
    f"{len(MATCH)} yamesasisha majibu ya awali na {120-len(MATCH)} ni mapya. "
    "Tafsiri mpya ambazo hazipo kwenye chanzo zinahitaji uhakiki. Hati hii ni rejea ya kazi, si tamko rasmi."
)
by_topic = {t["id"]: [] for t in kb["topics"]}
for faq in kb["faqs"]:
    by_topic.setdefault(faq["topic"], []).append(faq)
for t in kb["topics"]:
    faqs = by_topic.get(t["id"], [])
    review.add_heading(f"{t['title']['sw'].upper()} ({len(faqs)})", level=1)
    for faq in faqs:
        review.add_heading(faq["q"].get("sw") or faq["q"]["en"], level=2)
        source_update = faq.get("source_update")
        if source_update:
            number = source_update["question_number"]
            review.add_paragraph(f"Nyaraka ya leo, swali {number}: {parsed[number-1][2]}")
            review.add_paragraph("Hali ya tafsiri: inahitaji uhakiki wa lugha nyingine.")
        else:
            review.add_paragraph("SW: " + faq["answer"]["sw"].get("text", ""))
            review.add_paragraph("EN: " + faq["answer"]["en"].get("text", ""))
review.save(OUTPUT)
print(f"Q&A total {len(kb['faqs'])}; updated {len(MATCH)}; added {120-len(MATCH)}; document {OUTPUT}")
